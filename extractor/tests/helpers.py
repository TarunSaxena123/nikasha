"""Test helpers: multipart bodies, JSON patch application, documents built in memory."""

from __future__ import annotations

import copy
import io
import json
from pathlib import Path
from typing import Any

REPO = Path(__file__).resolve().parents[2]
SAMPLES = REPO / "samples"
CONTRACTS = REPO / "specs" / "contracts"
RECORDINGS = Path(__file__).parent / "recordings"


def multipart(data: bytes, filename: str = "my secret plan.pdf") -> tuple[bytes, str]:
    boundary = "----nikashatestboundary"
    body = (
        (
            f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\n'
            f"Content-Type: application/octet-stream\r\n\r\n"
        ).encode()
        + data
        + f"\r\n--{boundary}--\r\n".encode()
    )
    return body, f"multipart/form-data; boundary={boundary}"


def apply_patch(doc: dict[str, Any], ops: list[dict[str, Any]]) -> dict[str, Any]:
    out = copy.deepcopy(doc)
    for op in ops:
        parts = op["path"].split("/")[1:]
        cur: Any = out
        for i, seg in enumerate(parts[:-1]):
            nxt_index = parts[i + 1].isdigit()
            key: Any = int(seg) if isinstance(cur, list) else seg
            if isinstance(cur, list):
                while len(cur) <= key:
                    cur.append({})
                if not isinstance(cur[key], dict | list):
                    cur[key] = [] if nxt_index else {}
                cur = cur[key]
            else:
                if key not in cur or not isinstance(cur[key], dict | list):
                    cur[key] = [] if nxt_index else {}
                cur = cur[key]
        last: Any = parts[-1]
        if op["op"] == "remove":
            if isinstance(cur, list):
                cur.pop(int(last))
            else:
                cur.pop(last, None)
        elif isinstance(cur, list):
            idx = int(last)
            while len(cur) <= idx:
                cur.append(None)
            cur[idx] = copy.deepcopy(op["value"])
        else:
            cur[last] = copy.deepcopy(op["value"])
    return out


def answer_number(doc: dict[str, Any], question: dict[str, Any], value: float) -> dict[str, Any]:
    ops = [
        {"op": "add", "path": t["path"], "value": value * t.get("multiplier", 1)} for t in question["number"]["targets"]
    ]
    return apply_patch(doc, ops)


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def image_only_pdf() -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(100, 100, 200, 200, fill=1)
    c.showPage()
    c.save()
    return buf.getvalue()


def load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))
