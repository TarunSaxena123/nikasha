"""File → text, in memory only (REQ-UPL-04; feature 07 design §2.3, REQ-PRV-20).

The type is sniffed from the bytes, never from the file name (which is ignored). Nothing is written
to disk: pypdf and python-docx read from io.BytesIO.
"""

from __future__ import annotations

import io
import zipfile

from nikasha_extractor.core.errors import ExtractError

MAX_FILE_BYTES = 4 * 1024 * 1024
MAX_TEXT_CHARS = 20_000


def sniff(data: bytes) -> str:
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return "unknown"
        return "unknown"
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        return "unknown"
    printable = sum(ch.isprintable() or ch in "\r\n\t" for ch in text[:2000])
    return "txt" if text and printable / max(1, min(len(text), 2000)) > 0.95 else "unknown"


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtractError("UNREADABLE_FILE")
            except Exception as exc:  # noqa: BLE001 — any decryption failure means unreadable
                raise ExtractError("UNREADABLE_FILE") from exc
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise ExtractError("UNREADABLE_FILE") from exc


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 — corrupt package
        raise ExtractError("UNREADABLE_FILE") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_upload(data: bytes) -> str:
    """Plain text of an uploaded file. Raises UNSUPPORTED_FILE / UNREADABLE_FILE / PAYLOAD_TOO_LARGE."""
    if len(data) > MAX_FILE_BYTES:
        raise ExtractError("PAYLOAD_TOO_LARGE")
    kind = sniff(data)
    if kind == "pdf":
        text = _pdf_text(data)
    elif kind == "docx":
        text = _docx_text(data)
    elif kind == "txt":
        text = data.decode("utf-8")
    else:
        raise ExtractError("UNSUPPORTED_FILE")
    text = text.replace("\x00", "").strip()
    if not text:
        raise ExtractError("UNREADABLE_FILE")
    return text


def check_text(text: str) -> str:
    text = text.strip()
    if not text:
        raise ExtractError("BAD_REQUEST")
    if len(text) > MAX_TEXT_CHARS:
        raise ExtractError("PAYLOAD_TOO_LARGE")
    return text
